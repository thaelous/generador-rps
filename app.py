import io
import os
import re
import json
import time
import base64
import requests
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PILImage
import fitz  # PyMuPDF
import streamlit as st

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.pdfgen import canvas

st.set_page_config(page_title="Generador RPS AAM", page_icon="📄", layout="centered")
st.title("Generador Automático de RPS con Evidencias")
st.write("Sube la Factura, la Orden de Compra y las fotos de evidencia para generar tu Excel y PDF oficiales.")

api_key = st.secrets.get("GEMINI_API_KEY", os.getenv("GEMINI_API_KEY", ""))

if not api_key:
    api_key = st.text_input("Ingresa tu Gemini API Key:", type="password")

uploaded_factura = st.file_uploader("1. Factura en PDF (CFDI)", type=["pdf"])
uploaded_oc = st.file_uploader("2. Orden de Compra en PDF (OC)", type=["pdf"])
uploaded_fotos = st.file_uploader("3. Fotos de Evidencia ('DESPUÉS' - hasta 3 imágenes)", type=["png", "jpg", "jpeg"], accept_multiple_files=True)

SYSTEM_PROMPT = """
Eres un auditor contable corporativo. Se te proporcionan dos documentos:
1. Factura (CFDI en PDF).
2. Orden de Compra (OC en PDF).

Tu tarea:
1. Extraer los datos fiscales del CFDI y los generales de la OC para el formato RPS.
2. Identificar la página específica de la Orden de Compra (número de página base 1) donde aparece la partida o tabla de partidas facturadas.
3. Extraer todas las líneas facturadas. Si hay más de un concepto o partida, extraer cada uno en el arreglo 'lineas'. Para cada línea, identificar su número de línea en la OC y formatearlo estrictamente como 'X-1' (por ejemplo: '1-1', '2-1', '3-1').

Devuelve EXCLUSIVAMENTE un objeto JSON con la siguiente estructura:
{
  "orden_compra": string,
  "nombre_proveedor": string,
  "folio_factura": string,
  "tipo_servicio": string,
  "solicitante": string,
  "moneda": string,
  "subtotal": number,
  "pagina_oc_partida": number,
  "lineas": [
    {
      "linea_po": string,
      "cantidad": number,
      "unidad": string,
      "descripcion": string,
      "monto": number
    }
  ]
}

Reglas estrictas:
- 'pagina_oc_partida': número entero de la página del PDF de la OC donde está el renglón/partida facturada (ej. 4).
- 'lineas.linea_po': código en formato 'X-1' correspondiente a esa partida en la OC.
- 'lineas.descripcion': sólo el concepto del servicio, eliminando solicitante y número de OC/PO.
- Devuelve únicamente el JSON válido.
"""

def extraer_datos_con_oc(factura_bytes, oc_bytes, raw_key):
    clean_key = raw_key.strip().strip("'").strip('"')
    fac_b64 = base64.b64encode(factura_bytes).decode('utf-8')
    
    parts = [
        {"inline_data": {"mime_type": "application/pdf", "data": fac_b64}},
        {"text": "Factura emitida (CFDI)."}
    ]
    
    if oc_bytes:
        oc_b64 = base64.b64encode(oc_bytes).decode('utf-8')
        parts.extend([
            {"inline_data": {"mime_type": "application/pdf", "data": oc_b64}},
            {"text": "Orden de Compra oficial (OC). Cruza las partidas facturadas con las líneas de la OC, extrae cada partida con su línea 'X-1' e indica en 'pagina_oc_partida' la página exacta donde se visualiza el renglón."}
        ])
    else:
        parts.append({"text": "Extrae los datos únicamente de la factura."})

    payload = {
        "system_instruction": {"parts": [{"text": SYSTEM_PROMPT}]},
        "contents": [{"parts": parts}],
        "generationConfig": {
            "response_mime_type": "application/json",
            "temperature": 0.1
        }
    }
    
    headers = {
        "Content-Type": "application/json",
        "x-goog-api-key": clean_key
    }
    
    modelos = [
        "gemini-2.5-flash-lite",
        "gemini-2.5-flash",
        "gemini-3.6-flash"
    ]
    ultimo_error = None
    
    for mod in modelos:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{mod}:generateContent".strip()
        for intento in range(3):
            try:
                response = requests.post(url, headers=headers, json=payload, timeout=75)
                if response.status_code == 200:
                    data = response.json()
                    texto = data["candidates"][0]["content"]["parts"][0]["text"]
                    return json.loads(texto)
                elif response.status_code in (429, 503):
                    ultimo_error = f"Código {response.status_code}: Servidor ocupado ({mod}). Reintentando..."
                    time.sleep(2 * (intento + 1))
                    continue
                else:
                    ultimo_error = f"Código {response.status_code}: {response.text}"
                    break
            except Exception as e:
                ultimo_error = str(e)
                time.sleep(2)
                
    raise RuntimeError(ultimo_error)

def limpiar_descripcion(desc, solicitante, oc):
    texto = desc
    if solicitante and solicitante.strip():
        patron_sol = re.compile(re.escape(solicitante.strip()), re.IGNORECASE)
        texto = patron_sol.sub("", texto)
    texto = re.sub(r'(?i)(solicita(?:nte)?|atenci[oó]n|contacto)\s*[:\-]?\s*', '', texto)
    if oc and str(oc).strip():
        patron_oc = re.compile(rf'(?i)(?:oc|orden\s*(?:de)?\s*compra|po)\s*[:#\-]?\s*{re.escape(str(oc).strip())}')
        texto = patron_oc.sub("", texto)
        texto = re.sub(rf'\b{re.escape(str(oc).strip())}\b', '', texto)
    texto = re.sub(r'(?i)\b(?:oc|po)\s*[:#\-]?\b', '', texto)
    texto = re.sub(r'\s+', ' ', texto)
    texto = re.sub(r'^[\s,\.\-_/]+|[\s,\.\-_/]+$', '', texto)
    return texto.strip()

def extraer_pagina_completa_oc(oc_bytes, pagina_sugerida=None):
    doc = fitz.open(stream=oc_bytes, filetype="pdf")
    total_paginas = len(doc)
    target_idx = 0
    
    if pagina_sugerida and isinstance(pagina_sugerida, int) and 1 <= pagina_sugerida <= total_paginas:
        target_idx = pagina_sugerida - 1
    else:
        for i, page in enumerate(doc):
            t = page.get_text()
            if "Partida/Descripción" in t or "Precio ampliado" in t or "Precio unitario" in t:
                target_idx = i
                break
                
    page = doc[target_idx]
    pix = page.get_pixmap(dpi=150)
    img = PILImage.open(io.BytesIO(pix.tobytes("png")))
    
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format='PNG')
    img_byte_arr.seek(0)
    return img_byte_arr

def llenar_plantilla_excel(datos, oc_bytes=None, fotos_bytes=[], plantilla_path="plantilla_RPS.xlsx"):
    wb = openpyxl.load_workbook(plantilla_path)
    ws = wb["3151"] if "3151" in wb.sheetnames else wb.active
    
    folio = str(datos.get("folio_factura", "RPS"))
    oc_val = datos.get("orden_compra", "")
    solicitante_val = datos.get("solicitante", "")
    ws.title = folio
    
    if oc_val:
        ws.cell(row=7, column=5, value=int(oc_val) if str(oc_val).isdigit() else str(oc_val))
    
    ws.cell(row=9, column=5, value=datos.get("nombre_proveedor", ""))
    ws.cell(row=11, column=5, value=int(folio) if folio.isdigit() else str(folio))
    ws.cell(row=11, column=10, value=datos.get("tipo_servicio", "Entrenamiento"))
    
    lineas = datos.get("lineas", [])
    fila_inicio = 15
    for idx, l in enumerate(lineas):
        r = fila_inicio + idx
        if r > 24:
            break
        ws.cell(row=r, column=1, value=l.get("linea_po", "1-1"))
        ws.cell(row=r, column=2, value=l.get("cantidad", 1))
        
        unidad = str(l.get("unidad", "LOT")).strip()
        if "E48" in unidad.upper() or not unidad:
            unidad = "LOT"
        else:
            unidad = unidad.replace("E48", "").replace("-", "").strip()
        ws.cell(row=r, column=4, value=unidad)
        
        ws.cell(row=r, column=5, value=l.get("monto", 0))
        ws.cell(row=r, column=6, value=datos.get("moneda", "MXN"))
        
        desc_limpia = limpiar_descripcion(l.get("descripcion", ""), solicitante_val, oc_val)
        ws.cell(row=r, column=7, value=desc_limpia)
        ws.cell(row=r, column=16, value="YES")
        
    subtotal = datos.get("subtotal", 0)
    ws.cell(row=25, column=5, value=subtotal)
    ws.cell(row=25, column=6, value=subtotal)
    
    if solicitante_val:
        ws.cell(row=27, column=3, value=solicitante_val)
        
    ws._images.clear()

    # 1. ANTES: Más grande y centrado en T10 (430px x 570px)
    if oc_bytes:
        try:
            pag_oc = datos.get("pagina_oc_partida")
            img_oc_bytes = extraer_pagina_completa_oc(oc_bytes, pag_oc)
            img_oc = OpenpyxlImage(img_oc_bytes)
            img_oc.width = 430
            img_oc.height = 570
            ws.add_image(img_oc, "T10")
        except Exception:
            pass
            
    # 2. DESPUÉS: Un poco más grandes (hasta 150px de alto x 195px de ancho) en AD10, AD17, AD24
    celdas_despues = ["AD10", "AD17", "AD24"]
    for i, f_bytes in enumerate(fotos_bytes[:3]):
        try:
            pil_temp = PILImage.open(io.BytesIO(f_bytes))
            w_orig, h_orig = pil_temp.size
            
            # Altura deseada aumentada a 150px
            max_h = 150
            ratio = max_h / float(h_orig)
            w_calc = int(w_orig * ratio)
            
            # Ancho límite para no salirse de las columnas
            if w_calc > 195:
                w_calc = 195
                max_h = int(h_orig * (195 / float(w_orig)))

            b_arr = io.BytesIO(f_bytes)
            excel_img = OpenpyxlImage(b_arr)
            excel_img.width = w_calc
            excel_img.height = max_h
            ws.add_image(excel_img, celdas_despues[i])
        except Exception:
            pass
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generar_pdf_oficial(datos, oc_bytes=None, fotos_bytes=[]):
    buffer = io.BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    ancho, alto = letter
    
    # ------------------ PÁGINA 1: FORMATO RPS ------------------
    p.setFillColor(colors.HexColor("#C00000"))
    p.setFont("Helvetica-Bold", 16)
    p.drawString(40, alto - 45, "AAM")
    p.setFillColor(colors.black)
    p.setFont("Helvetica-Bold", 12)
    p.drawString(95, alto - 45, "REPLACEMENT PACKING SLIP")
    
    p.setFont("Helvetica-Bold", 6.5)
    p.setFillColor(colors.HexColor("#C00000"))
    p.drawString(40, alto - 58, "INSTRUCCIONES: TODOS LOS CAMPOS SOMBREADOS DEBEN SER LLENADOS PARA PODER PROCEDER CON EL RECIBO.")
    p.drawString(40, alto - 67, "ES NECESARIO ADJUNTAR LA FACTURA A ESTE DOCUMENTO")
    
    p.setFont("Helvetica-Bold", 8)
    p.setFillColor(colors.black)
    p.drawString(40, alto - 90, "ORDEN DE COMPRA #:")
    p.drawString(40, alto - 108, "NOMBRE DE PROVEEDOR:")
    p.drawString(40, alto - 126, "FACTURA #:")
    
    p.setFont("Helvetica", 8)
    oc_val = str(datos.get("orden_compra", ""))
    folio_val = str(datos.get("folio_factura", ""))
    proveedor_val = str(datos.get("nombre_proveedor", ""))
    solicitante_val = str(datos.get("solicitante", ""))
    tipo_servicio_val = str(datos.get("tipo_servicio", "Entrenamiento"))
    
    p.drawString(170, alto - 90, oc_val)
    p.drawString(170, alto - 108, proveedor_val)
    p.drawString(170, alto - 126, folio_val)
    
    p.setFont("Helvetica-Bold", 8)
    p.drawString(370, alto - 126, "TIPO DE SERVICIO A RECIBIR:")
    p.setFont("Helvetica", 8)
    p.drawString(500, alto - 126, tipo_servicio_val)
    
    y_tabla = alto - 150
    lineas = datos.get("lineas", [])
    num_filas = max(len(lineas), 1)
    altura_tabla = 20 + (num_filas * 18)
    
    p.setFillColor(colors.HexColor("#FFF2CC"))
    p.rect(40, y_tabla - altura_tabla, ancho - 80, altura_tabla, fill=1, stroke=1)
    
    p.setFillColor(colors.black)
    p.setFont("Helvetica-Bold", 7)
    p.drawString(45, y_tabla - 12, "# DE LINEA PO")
    p.drawString(110, y_tabla - 12, "CANT.")
    p.drawString(140, y_tabla - 12, "UOM")
    p.drawString(170, y_tabla - 12, "MONTO A RECIBIR")
    p.drawString(255, y_tabla - 12, "MONEDA")
    p.drawString(305, y_tabla - 12, "DESCRIPCIÓN DE LA LINEA A RECIBIR")
    p.drawString(525, y_tabla - 12, "COMPLETA?")
    
    subtotal = datos.get("subtotal", 0)
    moneda = datos.get("moneda", "MXN")
    
    p.setFont("Helvetica", 7)
    for i, l in enumerate(lineas):
        y_fila = y_tabla - 28 - (i * 18)
        desc_l = limpiar_descripcion(l.get("descripcion", ""), solicitante_val, oc_val)
        monto_l = l.get("monto", subtotal)
        p.drawString(45, y_fila, str(l.get("linea_po", "1-1")))
        p.drawString(115, y_fila, str(l.get("cantidad", 1)))
        p.drawString(140, y_fila, str(l.get("unidad", "LOT")))
        p.drawString(170, y_fila, f"${monto_l:,.2f}")
        p.drawString(260, y_fila, moneda)
        p.drawString(305, y_fila, desc_l[:48])
        p.drawString(535, y_fila, "YES")
        
    y_totales = y_tabla - altura_tabla - 25
    p.setFont("Helvetica-Bold", 8)
    p.drawString(40, y_totales, "MONTO TOTAL A RECIBIR:")
    p.drawString(170, y_totales, f"${subtotal:,.2f} {moneda}")
    
    p.drawString(40, y_totales - 30, "SOLICITANTE:")
    p.setFont("Helvetica", 8)
    p.drawString(120, y_totales - 30, solicitante_val)
    
    p.setFont("Helvetica-Bold", 8)
    p.drawString(40, y_totales - 55, "APROBADOR:")
    p.setFont("Helvetica", 8)
    p.drawString(120, y_totales - 55, "Laura Maciel Hernández 014036")
    
    p.showPage()
    
    # ------------------ PÁGINA 2: EVIDENCIAS ------------------
    p.setFillColor(colors.HexColor("#C00000"))
    p.setFont("Helvetica-Bold", 16)
    p.drawString(40, alto - 45, "AAM")
    p.setFillColor(colors.black)
    p.setFont("Helvetica-Bold", 12)
    p.drawString(95, alto - 45, "REPLACEMENT PACKING SLIP EVIDENCE")
    
    p.setFont("Helvetica-Bold", 6.5)
    p.setFillColor(colors.HexColor("#C00000"))
    p.drawString(40, alto - 58, "INSTRUCCIONES: LAS FOTOS DEBEN ESTAR DEL TAMAÑO DEL RECUADRO MARCADO Y RESOLUCIÓN DE CALIDAD")
    p.drawString(40, alto - 67, "NOTA: SOLO APLICA PARA SERVICIOS")
    
    p.setFont("Helvetica-Bold", 9)
    p.drawString(135, alto - 90, "FOTOS DEL ANTES")
    p.drawString(395, alto - 90, "FOTOS DEL DESPUES")
    
    w_box = 245
    h_box = 490
    y_box = alto - 595
    p.setStrokeColor(colors.HexColor("#A6A6A6"))
    p.setLineWidth(1)
    p.rect(40, y_box, w_box, h_box, fill=0, stroke=1)
    p.rect(320, y_box, w_box, h_box, fill=0, stroke=1)
    
    if oc_bytes:
        try:
            pag_oc = datos.get("pagina_oc_partida")
            img_oc_io = extraer_pagina_completa_oc(oc_bytes, pag_oc)
            pil_oc = PILImage.open(img_oc_io)
            temp_oc_path = f"/tmp/oc_full_{int(time.time()*1000)}.png"
            pil_oc.save(temp_oc_path)
            p.drawImage(temp_oc_path, 45, y_box + 10, width=w_box - 10, height=h_box - 20, preserveAspectRatio=True)
            if os.path.exists(temp_oc_path):
                os.remove(temp_oc_path)
        except Exception:
            pass
            
    if fotos_bytes:
        n_fotos = min(len(fotos_bytes), 3)
        h_disponible_por_foto = (h_box - 20) / n_fotos
        for i, fb in enumerate(fotos_bytes[:3]):
            try:
                p_foto = PILImage.open(io.BytesIO(fb))
                w_orig, h_orig = p_foto.size
                ratio = min(150 / w_orig, (h_disponible_por_foto - 10) / h_orig)
                w_render = w_orig * ratio
                h_render = h_orig * ratio
                
                temp_foto_path = f"/tmp/foto_desp_{i}_{int(time.time()*1000)}.png"
                p_foto.save(temp_foto_path)
                
                offset_y = (y_box + h_box - 10) - ((i + 1) * h_disponible_por_foto) + ((h_disponible_por_foto - h_render) / 2)
                offset_x = 320 + ((w_box - w_render) / 2)
                
                p.drawImage(temp_foto_path, offset_x, offset_y, width=w_render, height=h_render, preserveAspectRatio=True)
                if os.path.exists(temp_foto_path):
                    os.remove(temp_foto_path)
            except Exception:
                pass
                
    p.save()
    buffer.seek(0)
    return buffer

# ----------------- GESTIÓN DE SESIÓN PERSISTENTE -----------------
if "procesado" not in st.session_state:
    st.session_state.procesado = False
    st.session_state.excel_salida = None
    st.session_state.pdf_salida = None
    st.session_state.nombre_base = ""
    st.session_state.datos = None

if uploaded_factura and api_key:
    if st.button("Procesar y Generar Documentos"):
        with st.spinner("Analizando documentos con IA y ensamblando archivos..."):
            try:
                oc_bytes = uploaded_oc.getvalue() if uploaded_oc else None
                fotos_bytes = [f.getvalue() for f in uploaded_fotos] if uploaded_fotos else []
                
                datos = extraer_datos_con_oc(uploaded_factura.getvalue(), oc_bytes, api_key)
                
                excel_salida = llenar_plantilla_excel(
                    datos, 
                    oc_bytes=oc_bytes,
                    fotos_bytes=fotos_bytes
                )
                
                pdf_salida = generar_pdf_oficial(
                    datos,
                    oc_bytes=oc_bytes,
                    fotos_bytes=fotos_bytes
                )
                
                folio = str(datos.get("folio_factura", "RPS"))
                oc = str(datos.get("orden_compra", "OC"))
                
                st.session_state.procesado = True
                st.session_state.excel_salida = excel_salida.getvalue()
                st.session_state.pdf_salida = pdf_salida.getvalue()
                st.session_state.nombre_base = f"RPS {oc} {folio}"
                st.session_state.datos = datos
                
            except Exception as e:
                st.error(f"Error al procesar: {e}")

if st.session_state.procesado:
    total_lineas = len(st.session_state.datos.get("lineas", []))
    st.success(f"¡Documentos listos! ({total_lineas} partida(s) procesada(s))")
    
    with st.expander("Ver detalle de datos extraídos"):
        st.json(st.session_state.datos)
        
    col1, col2 = st.columns(2)
    with col1:
        st.download_button(
            label=f"📊 Descargar {st.session_state.nombre_base}.xlsx",
            data=st.session_state.excel_salida,
            file_name=f"{st.session_state.nombre_base}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="btn_dl_excel"
        )
    with col2:
        st.download_button(
            label=f"📄 Descargar {st.session_state.nombre_base}.pdf",
            data=st.session_state.pdf_salida,
            file_name=f"{st.session_state.nombre_base}.pdf",
            mime="application/pdf",
            key="btn_dl_pdf"
        )
